#!/usr/bin/env python3
"""
High-precision visual renderer and automated UI/UX auditor for:
Zenith_V_EV_Scooter_Launch_J_GANESH_KUMAR_REDDY_23BDS024.pptx
Renders each slide at 2000x1125 PNG with exact run-level word wrapping,
rounded rectangles, vertical centering, and strict overflow detection.
"""
import os
import io
import fitz
from pptx import Presentation
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE_TYPE, MSO_SHAPE
from PIL import Image, ImageDraw, ImageFont

FONT_REG = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
FONT_BOLD = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"

SCALE = 150.0


def get_font(size_pt, bold=False):
    px = max(10, int(round(size_pt * SCALE / 72.0 * 0.93)))
    path = FONT_BOLD if bold else FONT_REG
    try:
        return ImageFont.truetype(path, px)
    except Exception:
        return ImageFont.load_default()


def color_to_rgb(color_format, default=(30, 41, 59)):
    try:
        if color_format and color_format.rgb:
            rgb = color_format.rgb
            return (rgb[0], rgb[1], rgb[2])
    except Exception:
        pass
    return default


def fill_to_rgb(fill, default=None):
    try:
        if fill.type == 1:
            rgb = fill.fore_color.rgb
            return (rgb[0], rgb[1], rgb[2])
    except Exception:
        pass
    return default


def line_to_rgb(line, default=None):
    try:
        if line and line.fill and line.fill.type == 1:
            rgb = line.fill.fore_color.rgb
            return (rgb[0], rgb[1], rgb[2])
        if line and line.color and line.color.rgb:
            rgb = line.color.rgb
            return (rgb[0], rgb[1], rgb[2])
    except Exception:
        pass
    return default


def layout_and_draw_textframe(draw, tf, box_x, box_y, box_w, box_h, default_rgb=(30, 41, 59), force_vcenter=False, issues=None, label=""):
    if not tf or not tf.text.strip():
        return

    pad_l = (tf.margin_left.inches if tf.margin_left is not None else 0.05) * SCALE
    pad_r = (tf.margin_right.inches if tf.margin_right is not None else 0.05) * SCALE
    pad_t = (tf.margin_top.inches if tf.margin_top is not None else 0.03) * SCALE
    pad_b = (tf.margin_bottom.inches if tf.margin_bottom is not None else 0.03) * SCALE

    max_w = max(20, box_w - pad_l - pad_r)
    para_layouts = []

    for p in tf.paragraphs:
        if not p.text.strip():
            para_layouts.append({"lines": [], "space_after": 2, "align": PP_ALIGN.LEFT, "line_h": 8})
            continue

        p_size = p.font.size.pt if p.font and p.font.size else 10.5
        p_bold = bool(p.font.bold) if p.font and p.font.bold is not None else False
        p_rgb = color_to_rgb(p.font.color if p.font else None, default_rgb)

        tokens = []
        if p.runs:
            for r in p.runs:
                r_size = r.font.size.pt if r.font and r.font.size else p_size
                r_bold = bool(r.font.bold) if r.font and r.font.bold is not None else p_bold
                r_rgb = color_to_rgb(r.font.color if r.font else None, p_rgb)
                r_font = get_font(r_size, r_bold)
                sublines = r.text.split("\n")
                for s_i, subline in enumerate(sublines):
                    if s_i > 0:
                        tokens.append(("\n", r_font, r_rgb, r_size))
                    parts = subline.split(" ")
                    for idx, part in enumerate(parts):
                        text_piece = part if idx == len(parts) - 1 else part + " "
                        if text_piece:
                            tokens.append((text_piece, r_font, r_rgb, r_size))
        else:
            r_font = get_font(p_size, p_bold)
            sublines = p.text.split("\n")
            for s_i, subline in enumerate(sublines):
                if s_i > 0:
                    tokens.append(("\n", r_font, p_rgb, p_size))
                parts = subline.split(" ")
                for idx, part in enumerate(parts):
                    text_piece = part if idx == len(parts) - 1 else part + " "
                    if text_piece:
                        tokens.append((text_piece, r_font, p_rgb, p_size))

        lines = []
        curr_line = []
        curr_w = 0.0
        max_pt_in_para = p_size

        for piece, fnt, rgb, pt in tokens:
            max_pt_in_para = max(max_pt_in_para, pt)
            if piece == "\n":
                lines.append(curr_line)
                curr_line = []
                curr_w = 0.0
                continue
            bbox = draw.textbbox((0, 0), piece, font=fnt)
            w = bbox[2] - bbox[0]
            if curr_line and (curr_w + w > max_w):
                lines.append(curr_line)
                stripped = piece.lstrip(" ")
                if stripped:
                    bbox2 = draw.textbbox((0, 0), stripped, font=fnt)
                    curr_line = [(stripped, fnt, rgb)]
                    curr_w = bbox2[2] - bbox2[0]
                else:
                    curr_line = []
                    curr_w = 0.0
            else:
                curr_line.append((piece, fnt, rgb))
                curr_w += w
        if curr_line:
            lines.append(curr_line)

        line_h = max(11, int(round(max_pt_in_para * SCALE / 72.0 * 1.18)))
        space_after = (p.space_after.pt * SCALE / 72.0) if p.space_after else 1.5
        align = p.alignment or PP_ALIGN.LEFT
        para_layouts.append(
            {
                "lines": lines,
                "space_after": space_after,
                "align": align,
                "line_h": line_h,
            }
        )

    total_h = 0.0
    for idx_pl, pl in enumerate(para_layouts):
        if not pl["lines"]:
            total_h += pl["line_h"] * 0.40
        else:
            add_sa = pl["space_after"] if idx_pl < len(para_layouts) - 1 else 0.0
            total_h += len(pl["lines"]) * pl["line_h"] + add_sa

    if issues is not None and total_h > box_h + 4:
        snippet = tf.text.strip().replace("\n", " ")[:40]
        issues.append(f"Overflow ({total_h:.0f}px > {box_h:.0f}px) in {label}: '{snippet}...'")

    v_anchor = tf.vertical_anchor
    if force_vcenter or v_anchor == MSO_ANCHOR.MIDDLE:
        curr_y = box_y + max(2.0, (box_h - total_h) / 2.0)
    else:
        curr_y = box_y + pad_t

    for idx_pl, pl in enumerate(para_layouts):
        if not pl["lines"]:
            curr_y += pl["line_h"] * 0.40
            continue
        for line in pl["lines"]:
            if curr_y + pl["line_h"] * 0.78 > box_y + box_h + 6:
                break
            line_w = 0.0
            for piece, fnt, _ in line:
                bb = draw.textbbox((0, 0), piece, font=fnt)
                line_w += bb[2] - bb[0]
            if pl["align"] == PP_ALIGN.CENTER:
                tx = box_x + pad_l + max(0.0, (max_w - line_w) / 2.0)
            elif pl["align"] == PP_ALIGN.RIGHT:
                tx = box_x + box_w - pad_r - line_w
            else:
                tx = box_x + pad_l

            for piece, fnt, rgb in line:
                draw.text((tx, curr_y), piece, font=fnt, fill=rgb)
                bb = draw.textbbox((0, 0), piece, font=fnt)
                tx += bb[2] - bb[0]
            curr_y += pl["line_h"]
        if idx_pl < len(para_layouts) - 1:
            curr_y += pl["space_after"]


def render_and_audit(pptx_path, out_dir="assets/slides_preview", pdf_out="Zenith_V_EV_Scooter_Launch_J_GANESH_KUMAR_REDDY_23BDS024.pdf"):
    os.makedirs(out_dir, exist_ok=True)
    prs = Presentation(pptx_path)
    W_in = prs.slide_width.inches
    H_in = prs.slide_height.inches
    W_px = int(round(W_in * SCALE))
    H_px = int(round(H_in * SCALE))

    print(f"Presentation: {pptx_path} ({len(prs.slides)} slides, {W_in:.3f}x{H_in:.2f} in)")
    png_paths = []

    for idx, slide in enumerate(prs.slides, start=1):
        img = Image.new("RGB", (W_px, H_px), (248, 250, 252))
        draw = ImageDraw.Draw(img)

        issues = []
        id_mentions = 0

        for shp in slide.shapes:
            x = shp.left.inches * SCALE
            y = shp.top.inches * SCALE
            w = shp.width.inches * SCALE
            h = shp.height.inches * SCALE

            if shp.left.inches < -0.05 or shp.top.inches < -0.05:
                issues.append(f"Shape '{shp.name}' negative coords")
            if shp.left.inches + shp.width.inches > W_in + 0.08:
                issues.append(f"Shape '{shp.name}' exceeds right bound")
            if shp.top.inches + shp.height.inches > H_in + 0.08:
                issues.append(f"Shape '{shp.name}' exceeds bottom bound")

            if shp.has_text_frame and "J GANESH KUMAR REDDY" in shp.text_frame.text:
                id_mentions += 1

            if shp.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    pic_im = Image.open(io.BytesIO(shp.image.blob)).convert("RGBA")
                    pic_im = pic_im.resize((max(1, int(round(w))), max(1, int(round(h)))), Image.LANCZOS)
                    img.paste(pic_im, (int(round(x)), int(round(y))), pic_im)
                except Exception as e:
                    issues.append(f"Image error: {e}")
                continue

            if shp.has_table:
                tbl = shp.table
                n_rows = len(tbl.rows)
                col_widths = [c.width.inches * SCALE for c in tbl.columns]
                row_heights = [r.height.inches * SCALE for r in tbl.rows]
                total_rh = sum(row_heights)
                if total_rh > 0:
                    row_heights = [rh * (h / total_rh) for rh in row_heights]

                cy = y
                for r_i in range(n_rows):
                    cx = x
                    rh = row_heights[r_i]
                    for c_i, cw in enumerate(col_widths):
                        cell = tbl.cell(r_i, c_i)
                        bg = fill_to_rgb(cell.fill, (255, 255, 255))
                        draw.rectangle([cx, cy, cx + cw, cy + rh], fill=bg, outline=(203, 213, 225), width=1)
                        if cell.text_frame:
                            layout_and_draw_textframe(
                                draw,
                                cell.text_frame,
                                cx,
                                cy,
                                cw,
                                rh,
                                default_rgb=(30, 41, 59),
                                force_vcenter=True,
                                issues=issues,
                                label=f"Table({r_i},{c_i})",
                            )
                        cx += cw
                    cy += rh
                continue

            fill_rgb = fill_to_rgb(shp.fill, None)
            line_rgb = line_to_rgb(shp.line, None)
            line_w = max(1, int(round(shp.line.width.pt * SCALE / 72.0))) if (shp.line and shp.line.width) else 1

            is_rounded = False
            try:
                if shp.auto_shape_type == MSO_SHAPE.ROUNDED_RECTANGLE:
                    is_rounded = True
            except Exception:
                pass

            if fill_rgb is not None or line_rgb is not None:
                if is_rounded:
                    radius = min(12, int(min(w, h) * 0.20))
                    draw.rounded_rectangle(
                        [x, y, x + w, y + h],
                        radius=radius,
                        fill=fill_rgb,
                        outline=line_rgb,
                        width=line_w if line_rgb else 0,
                    )
                else:
                    draw.rectangle(
                        [x, y, x + w, y + h],
                        fill=fill_rgb,
                        outline=line_rgb,
                        width=line_w if line_rgb else 0,
                    )

            if shp.has_text_frame:
                vcenter = is_rounded and (h < 0.55 * SCALE)
                layout_and_draw_textframe(
                    draw,
                    shp.text_frame,
                    x,
                    y,
                    w,
                    h,
                    force_vcenter=vcenter,
                    issues=issues,
                    label=shp.name,
                )

        notes_txt = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
        if "J GANESH KUMAR REDDY 23BDS024" in notes_txt:
            id_mentions += 1

        out_png = os.path.join(out_dir, f"slide_{idx}.png")
        img.save(out_png)
        png_paths.append(out_png)
        status = "OK" if (not issues and id_mentions >= 2) else f"WARN: {issues}"
        print(f"  Slide {idx}: Student ID occurrences = {id_mentions} | Verified {status}")

    doc = fitz.open()
    for p in png_paths:
        page = doc.new_page(width=960, height=540)
        page.insert_image(fitz.Rect(0, 0, 960, 540), filename=p)
    doc.save(pdf_out)
    doc.close()
    print(f"Saved PDF preview: {pdf_out}")


if __name__ == "__main__":
    render_and_audit("Zenith_V_EV_Scooter_Launch_J_GANESH_KUMAR_REDDY_23BDS024.pptx")
