"""Design system + text/layout engine — light academic theme (Calibri)."""
import os, re
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from PIL import ImageFont

EMU_IN = 914400
def I(v): return Emu(int(round(v * EMU_IN)))

# ─────────────────────────────────────────── palette
NAVY   = RGBColor(0x12, 0x26, 0x3F)
INK    = RGBColor(0x1B, 0x2A, 0x41)
BODY   = RGBColor(0x44, 0x54, 0x6A)
MUTED  = RGBColor(0x7B, 0x8B, 0x9E)
LINE   = RGBColor(0xDC, 0xE3, 0xEC)
LINE_D = RGBColor(0xC3, 0xCE, 0xDC)
BLUE   = RGBColor(0x2B, 0x6C, 0xB0)
BLUE_L = RGBColor(0xEC, 0xF2, 0xFA)
BLUE_M = RGBColor(0xC9, 0xDC, 0xF0)
TEAL   = RGBColor(0x0F, 0x76, 0x6E)
TEAL_L = RGBColor(0xE9, 0xF5, 0xF2)
AMBER  = RGBColor(0xB4, 0x53, 0x09)
AMBER_L= RGBColor(0xFD, 0xF4, 0xE7)
RED    = RGBColor(0xB3, 0x26, 0x1E)
RED_L  = RGBColor(0xFB, 0xED, 0xEC)
GREEN  = RGBColor(0x2F, 0x6B, 0x4F)
GREY_L = RGBColor(0xF7, 0xF9, 0xFC)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)

FONT = "Calibri"
FDIR = "/home/user/.local/share/fonts"
_FILES = {400: "Carlito-Regular.ttf", 500: "Carlito-Regular.ttf",
          600: "Carlito-Bold.ttf", 700: "Carlito-Bold.ttf", 800: "Carlito-Bold.ttf"}
_cache = {}

def _pil(size_pt, weight):
    key = (round(size_pt * 4), weight)
    if key not in _cache:
        f = _FILES[400] if weight < 600 else _FILES[700]
        _cache[key] = ImageFont.truetype(os.path.join(FDIR, f), int(round(size_pt * 4)))
    return _cache[key]

def text_w(s, size_pt, weight=400):
    if not s: return 0.0
    return _pil(size_pt, weight).getlength(s) / 4.0

# ─────────────────────────────────────────── measuring
def _runs_of(p):
    return p["runs"] if "runs" in p else [(p.get("text", ""), p)]

def measure(paras, width_in):
    maxw = width_in * 72.0
    total = 0.0
    for p in paras:
        size, lm = p.get("size", 10), p.get("line", 1.22)
        tokens = []
        for txt, ro in _runs_of(p):
            sz, wt = ro.get("size", size), ro.get("weight", p.get("weight", 400))
            for tok in re.split(r"(\s+)", txt):
                if tok:
                    tokens.append((tok, text_w(tok, sz, wt), tok.isspace()))
        lines, cur = 1, 0.0
        for tok, w, is_space in tokens:
            if is_space:
                if cur > 0: cur += w
                continue
            if cur + w > maxw and cur > 0:
                lines += 1; cur = w
            else:
                cur += w
        total += lines * size * lm * 1.02 + p.get("sb", 0) + p.get("sa", 0)
    return total / 72.0

# ─────────────────────────────────────────── primitives
def no_shadow(shape):
    try: shape.shadow.inherit = False
    except Exception: pass

def rect(slide, x, y, w, h, fill=None, line=None, lw=0.75, radius=None, shape=None):
    if shape is None:
        shape = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(shape, I(x), I(y), I(w), I(h))
    no_shadow(s)
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try: s.adjustments[0] = radius
        except Exception: pass
    if fill is None: s.fill.background()
    else:
        s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None: s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(lw)
    s.text_frame.word_wrap = True
    return s

def hline(slide, x, y, w, color=LINE, lw=0.75):
    c = slide.shapes.add_connector(1, I(x), I(y), I(x + w), I(y))
    c.line.color.rgb = color; c.line.width = Pt(lw); return c

def vline(slide, x, y, h, color=LINE, lw=0.75):
    c = slide.shapes.add_connector(1, I(x), I(y), I(x), I(y + h))
    c.line.color.rgb = color; c.line.width = Pt(lw); return c

ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
ANCHOR = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}

def _style(r, size, color, weight, italic, spc, font):
    r.font.size = Pt(size)
    r.font.bold = weight >= 600
    r.font.italic = italic
    r.font.name = font
    r.font.color.rgb = color
    rPr = r._r.get_or_add_rPr()
    if spc: rPr.set("spc", str(int(spc * 100)))
    lat = rPr.find(qn("a:latin"))
    if lat is not None: lat.set("typeface", font)

def text(slide, x, y, w, h, paras, *, align="l", anchor="t", margins=(0, 0, 0, 0),
         shape=None, check=True, label=""):
    if isinstance(paras, str):
        paras = [{"text": paras}]
    tb = shape if shape is not None else slide.shapes.add_textbox(I(x), I(y), I(w), I(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left, tf.margin_right, tf.margin_top, tf.margin_bottom = [I(m) for m in margins]
    tf.vertical_anchor = ANCHOR[anchor]
    first = True
    for p in paras:
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        para.alignment = ALIGN[p.get("align", align)]
        para.line_spacing = p.get("line", 1.22)
        if p.get("sb"): para.space_before = Pt(p["sb"])
        if p.get("sa"): para.space_after = Pt(p["sa"])
        size, color, weight = p.get("size", 10), p.get("color", BODY), p.get("weight", 400)
        for txt_, ro in _runs_of(p):
            r = para.add_run(); r.text = txt_
            _style(r, ro.get("size", size), ro.get("color", color), ro.get("weight", weight),
                   ro.get("i", p.get("i", False)), ro.get("spc", p.get("spc", 0)), FONT)
    if check:
        need = measure(paras, w - margins[0] - margins[1])
        avail = h - margins[2] - margins[3]
        if need > avail + 0.015:
            print(f"  ! OVERFLOW [{label or paras[0].get('text','')[:40]!r}] "
                  f"need {need:.2f} > {avail:.2f}")
    return tb

def pill_w(label, size=8.2, weight=700, spc=1.0, pad=0.14):
    return text_w(label, size, weight) / 72.0 + spc * len(label) / 72.0 + 2 * pad

def pill(slide, x, y, label, fill, fg=WHITE, size=8.2, weight=700, spc=1.0, h=0.24,
         pad=0.14, border=None):
    w = pill_w(label, size, weight, spc, pad)
    s = rect(slide, x, y, w, h, fill=fill, line=border, lw=0.75, radius=0.5)
    text(slide, x, y, w, h, [{"text": label, "size": size, "color": fg, "weight": weight,
                              "spc": spc, "line": 1.0, "align": "c"}],
         align="c", anchor="m", shape=s)
    return w

def picture(slide, path, x, y, w=None, h=None, crop=None):
    from PIL import Image
    iw, ih = Image.open(path).size
    if w is None: w = h * iw / ih
    if h is None: h = w * ih / iw
    p = slide.shapes.add_picture(path, I(x), I(y), I(w), I(h))
    if crop:
        p.crop_left, p.crop_top, p.crop_right, p.crop_bottom = crop
    return p

# ─────────────────────────────────────────── page frame
PAGE_W, PAGE_H = 13.333, 7.5
MX = 0.72
CW = PAGE_W - 2 * MX

def new_slide(prs, bg=WHITE):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, PAGE_W, PAGE_H, fill=bg, shape=MSO_SHAPE.RECTANGLE)
    return s

def header(slide, section, title, sub=None, session=None, tag_color=BLUE,
           session_link=None):
    """Draws the standard content-slide header. Returns the y where content may start."""
    text(slide, MX, 0.34, 6.2, 0.20,
         [{"text": section.upper(), "size": 8.0, "color": tag_color, "weight": 700,
           "spc": 1.5, "line": 1.0}], label="section")
    if session:
        st = text(slide, PAGE_W - MX - 5.4, 0.34, 5.4, 0.20,
                  [{"text": session.upper(), "size": 7.6,
                    "color": BLUE if session_link else MUTED,
                    "weight": 600, "spc": 1.1, "line": 1.0, "align": "r"}],
                  align="r", label="session")
        if session_link:
            link_run(slide, st, session_link)
    th = measure([{"text": title, "size": 23, "weight": 700, "line": 1.12}], CW - 0.6)
    text(slide, MX, 0.58, CW, max(th, 0.42),
         [{"text": title, "size": 23, "color": NAVY, "weight": 700, "line": 1.12}], label="title")
    y = 0.58 + th + 0.10
    if sub:
        sh = measure([{"text": sub, "size": 10.2, "line": 1.30}], CW - 1.2)
        text(slide, MX, y, CW - 0.6, max(sh, 0.2),
             [{"text": sub, "size": 10.2, "color": BODY, "line": 1.30}], label="sub")
        y += sh + 0.10
    hline(slide, MX, y + 0.04, CW, LINE, 0.75)
    return y + 0.26

def footer(slide, page, total=9):
    hline(slide, MX, 6.98, CW, LINE, 0.75)
    text(slide, MX, 7.08, 1.46, 0.18,
         [{"text": "J GANESH KUMAR REDDY", "size": 8.2, "color": NAVY, "weight": 700,
           "spc": 0.4, "line": 1.0}], label="foot-name")
    text(slide, MX + 1.50, 7.08, 0.82, 0.18,
         [{"text": "·  23BDS024", "size": 8.2, "color": MUTED, "weight": 500, "line": 1.0}],
         label="foot-roll")
    text(slide, 4.95, 7.08, 3.42, 0.18,
         [{"text": "Managerial Economics  ·  Zenith-V Case Study  ·  IIIT Dharwad",
           "size": 7.6, "color": MUTED, "weight": 400, "line": 1.0, "align": "c"}],
         align="c", label="foot-course")
    text(slide, PAGE_W - MX - 1.10, 7.08, 1.10, 0.18,
         [{"text": f"Page {page} of {total}", "size": 7.8, "color": MUTED, "weight": 600,
           "line": 1.0, "align": "r"}], align="r", label="foot-page")


def link_run(slide, shape, url):
    """Attach a hyperlink to every run of a text box (keeps the visible styling)."""
    for para in shape.text_frame.paragraphs:
        for run in para.runs:
            try:
                run.hyperlink.address = url
            except Exception as exc:      # pragma: no cover
                print("  ! link failed:", exc)


def bottom_line(slide, runs_or_text, y=6.44, fill=BLUE_L, accent=BLUE, label="BOTTOM LINE"):
    """Full-width closing strip so every content slide ends on the same baseline."""
    if isinstance(runs_or_text, str):
        runs = [(runs_or_text, {"size": 9.0, "color": NAVY})]
    else:
        runs = runs_or_text
    h = 0.44
    rect(slide, MX, y, CW, h, fill=fill, line=None, radius=0.04)
    rect(slide, MX, y, 0.05, h, fill=accent, shape=MSO_SHAPE.RECTANGLE)
    text(slide, MX + 0.20, y, CW - 0.40, h,
         [{"runs": [(label + "   ", {"size": 7.6, "color": accent, "weight": 700, "spc": 1.0})]
                    + runs, "line": 1.24}], anchor="m", label="bottom-line")


def card(slide, x, y, w, h, fill=GREY_L, line=LINE, radius=0.05):
    return rect(slide, x, y, w, h, fill=fill, line=line, lw=0.75, radius=radius)

def bullets(slide, x, y, w, items, size=9.4, gap=0.075, color=BODY, marker="▪",
            marker_color=BLUE, line=1.26, label=""):
    """Simple bulleted list. items: list of str or list of (text, opts) runs."""
    yy = y
    for it in items:
        if isinstance(it, str):
            runs = [(it, {"size": size, "color": color})]
        else:
            runs = [(" ", {})] if False else it
        p = [{"runs": [(" " + " " if False else "", {})] + runs, "line": line, "size": size}]
        # marker on its own textbox for stable alignment
        text(slide, x, yy, 0.17, 0.2,
             [{"text": marker, "size": size - 0.6, "color": marker_color, "weight": 700,
               "line": 1.0}], label="bullet")
        body = [{"runs": runs, "size": size, "color": color, "line": line}]
        h = measure(body, w - 0.24)
        text(slide, x + 0.22, yy - 0.01, w - 0.22, h, body, label=label or "bullet-txt")
        yy += h + gap
    return yy
