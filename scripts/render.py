"""Approximate raster preview of the PPTX for layout QA (Calibri/Carlito metrics)."""
import os, sys, re, io
sys.path.insert(0, os.path.dirname(__file__))
from pptx import Presentation
from PIL import Image, ImageDraw
from deck_lib import _pil, text_w

PPI = 144.0
def px(v): return int(round(v * PPI))

def get_rgb(color, default=(30, 42, 58)):
    try:
        c = color.rgb
        return (c[0], c[1], c[2])
    except Exception:
        return default

def fill_kind(sh):
    from pptx.oxml.ns import qn
    try:
        spPr = sh._element.spPr
    except Exception:
        return None
    if spPr is None: return None
    if spPr.find(qn("a:gradFill")) is not None: return "grad"
    if spPr.find(qn("a:solidFill")) is not None: return "solid"
    return None

def render(path, outdir):
    prs = Presentation(path)
    W, H = px(prs.slide_width / 914400), px(prs.slide_height / 914400)
    os.makedirs(outdir, exist_ok=True)
    for idx, slide in enumerate(prs.slides, 1):
        img = Image.new("RGB", (W, H), (255, 255, 255))
        d = ImageDraw.Draw(img)
        for sh in slide.shapes:
            x, y = px(sh.left / 914400), px(sh.top / 914400)
            w, h = px(sh.width / 914400), px(sh.height / 914400)
            if sh.shape_type == 13:
                try:
                    im = Image.open(io.BytesIO(sh.image.blob)).convert("RGBA")
                    cl, ct, cr, cb = sh.crop_left, sh.crop_top, sh.crop_right, sh.crop_bottom
                    iw, ih = im.size
                    box = (int(cl * iw), int(ct * ih), int(iw * (1 - cr)), int(ih * (1 - cb)))
                    im = im.crop(box).resize((max(w, 1), max(h, 1)), Image.LANCZOS)
                    img.paste(im, (x, y), im)
                    d = ImageDraw.Draw(img)
                except Exception:
                    pass
                continue
            if sh.shape_type == 20:
                try:
                    col = get_rgb(sh.line.color, (200, 200, 200))
                    lw = max(1, int(round((sh.line.width.pt if sh.line.width else 0.75) * PPI / 72)))
                    d.line([x, y, x + w, y + h], fill=col, width=lw)
                except Exception:
                    pass
                continue
            kind = fill_kind(sh)
            try:
                adj = sh.adjustments[0]
            except Exception:
                adj = 0
            rad = int(adj * min(w, h)) if adj else 0
            if kind == "solid":
                rgb = get_rgb(sh.fill.fore_color, None)
                if rgb:
                    if sh.shape_type == 1 and False:
                        pass
                    if rad:
                        d.rounded_rectangle([x, y, x + w, y + h], radius=rad, fill=rgb)
                    elif getattr(sh, "auto_shape_type", None) is not None and str(sh.auto_shape_type).startswith("OVAL"):
                        d.ellipse([x, y, x + w, y + h], fill=rgb)
                    else:
                        d.rectangle([x, y, x + w, y + h], fill=rgb)
            try:
                if sh.line.fill.type == 1 and sh.line.width:
                    col = get_rgb(sh.line.color, (210, 218, 228))
                    lw = max(1, int(round(sh.line.width.pt * PPI / 72)))
                    if rad:
                        d.rounded_rectangle([x, y, x + w, y + h], radius=rad, outline=col, width=lw)
                    else:
                        d.rectangle([x, y, x + w, y + h], outline=col, width=lw)
            except Exception:
                pass
            if sh.has_text_frame:
                draw_tf(d, sh, x, y, w, h)
        img.save(os.path.join(outdir, f"p{idx}.png"))

def draw_tf(d, sh, x, y, w, h):
    tf = sh.text_frame
    ml = px(tf.margin_left / 914400) if tf.margin_left is not None else px(0.1)
    mr = px(tf.margin_right / 914400) if tf.margin_right is not None else px(0.1)
    mt = px(tf.margin_top / 914400) if tf.margin_top is not None else px(0.05)
    mb = px(tf.margin_bottom / 914400) if tf.margin_bottom is not None else px(0.05)
    inner_w = w - ml - mr
    if inner_w <= 4: return
    blocks = []
    for p in tf.paragraphs:
        runs = []
        for r in p.runs:
            rPr = r._r.find("{http://schemas.openxmlformats.org/drawingml/2006/main}rPr")
            size = r.font.size.pt if r.font.size else 18.0
            weight = 700 if r.font.bold else 400
            col = get_rgb(r.font.color, (30, 42, 58))
            spc = int(rPr.get("spc")) / 100.0 if (rPr is not None and rPr.get("spc")) else 0.0
            runs.append((r.text, size, weight, col, spc))
        if not runs:
            blocks.append(None); continue
        size = max(r[1] for r in runs)
        lst = p.line_spacing if p.line_spacing else 1.2
        sb = p.space_before.pt if p.space_before else 0
        al = str(p.alignment) if p.alignment is not None else "LEFT"
        blocks.append((size, lst, sb, al, runs))
    out_lines, total_h = [], 0.0
    for b in blocks:
        if b is None:
            out_lines.append(None); continue
        size, lst, sb, al, runs = b
        total_h += sb
        words = []
        for text_, sz, wt, col, spc in runs:
            for tok in re.split(r"(\s+)", text_):
                if tok == "":
                    continue
                tw = text_w(tok, sz, wt) * PPI / 72
                words.append((tok, sz, wt, col, spc, tw, tok.isspace()))
        line, lw = [], 0.0
        for tok, sz, wt, col, spc, tw, is_space in words:
            if is_space:
                if line:
                    line.append((tok, sz, wt, col, spc, tw))
                    lw += tw
                continue
            if lw + tw > inner_w and line:
                out_lines.append((line, size, lst, al, 0)); line, lw = [], 0.0
            line.append((tok, sz, wt, col, spc, tw)); lw += tw
        out_lines.append((line, size, lst, al, 0))
    heights = []
    for ln in out_lines:
        heights.append(8 if ln is None else ln[1] * ln[2] * 1.02 * PPI / 72)
    total = sum(heights)
    anchor = str(tf.vertical_anchor) if tf.vertical_anchor is not None else "TOP"
    ypos = y + mt
    if "MIDDLE" in anchor: ypos = y + mt + (h - mt - mb - total) / 2
    elif "BOTTOM" in anchor: ypos = y + h - mb - total
    for i, ln in enumerate(out_lines):
        if ln is None:
            ypos += 8; continue
        line, size, lst, al = ln[0], ln[1], ln[2], ln[3]
        lh = heights[i]
        baseline = ypos + size * PPI / 72 * 0.94
        widths = sum(t[-1] for t in line) + sum(t[4] * PPI / 72 for t in line)
        if "CENTER" in al: xx = x + ml + (inner_w - widths) / 2
        elif "RIGHT" in al: xx = x + ml + (inner_w - widths)
        else: xx = x + ml
        for tok, sz, wt, col, spc, tw in line:
            if not tok.isspace():
                f = _pil(sz, wt).font_variant(size=int(round(sz * PPI / 72)))
                d.text((xx, baseline), tok, font=f, fill=col, anchor="ls")
            xx += tw + spc * PPI / 72
        ypos += lh

if __name__ == "__main__":
    p = sys.argv[1] if len(sys.argv) > 1 else "../Zenith-V_Launch_Decision_JGaneshKumarReddy_23BDS024.pptx"
    o = sys.argv[2] if len(sys.argv) > 2 else "/tmp/preview"
    render(p, o)
    print("rendered to", o)
